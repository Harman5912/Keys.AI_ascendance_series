"""Independent artifact validation.

A worker's ``ok`` is a claim. This module checks the claim against the file on
disk: it exists, it is inside the task workspace, it is non-trivial, and — for
every format with a parser available — it reopens and reports real structure
(pages, slides, sheets, image dimensions, parsed rows).

Validation never edits a file. The service uses the result to decide whether a
task is COMPLETED or goes into REPAIRING/FAILED.
"""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path

from universal_agent import security


def _result(path: Path, ok: bool, detail: str, kind: str = "",
            warning: str = "") -> dict:
    return {"ok": bool(ok), "path": str(path), "kind": kind or path.suffix.lstrip("."),
            "detail": detail, "warning": warning}


def validate_artifact(path, workspace=None) -> dict:
    p = Path(str(path or ""))
    try:
        if workspace is not None:
            p = security.resolve_within(workspace, p)
        else:
            p = p.expanduser().resolve()
    except security.SecurityError as exc:
        return _result(p, False, str(exc))
    if not p.is_file():
        return _result(p, False, "file does not exist")
    size = p.stat().st_size
    if size == 0:
        return _result(p, False, "file is empty")
    suffix = p.suffix.lower()

    try:
        if suffix == ".docx":
            from docx import Document
            doc = Document(str(p))
            paragraphs = [x for x in doc.paragraphs if x.text.strip()]
            headings = [x.text.strip() for x in paragraphs
                        if x.style and x.style.name.lower().startswith("heading")]
            if not paragraphs:
                return _result(p, False, "Word document has no readable paragraphs")
            return _result(p, True,
                           f"{len(paragraphs)} paragraph(s), "
                           f"{len(headings)} heading(s), {len(doc.tables)} table(s)")
        if suffix == ".pdf":
            try:
                from pypdf import PdfReader
                reader = PdfReader(str(p))
                pages = len(reader.pages)
                text = (reader.pages[0].extract_text() or "") if pages else ""
                if pages < 1:
                    return _result(p, False, "PDF has no pages")
                return _result(p, True,
                               f"{pages} page(s), first page "
                               f"{len(text.strip())} chars of text")
            except ImportError:
                head = p.read_bytes()[:4096]
                if b"%PDF" not in head:
                    return _result(p, False, "file is not a PDF")
                return _result(p, True, "PDF signature valid (pypdf unavailable)",
                               warning="page count not verified")
        if suffix == ".pptx":
            from pptx import Presentation
            prs = Presentation(str(p))
            slides = len(prs.slides)
            if slides < 1:
                return _result(p, False, "presentation has no slides")
            texts = 0
            for slide in prs.slides:
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                        texts += 1
            if texts == 0:
                return _result(p, False, "presentation slides contain no text")
            return _result(p, True, f"{slides} slide(s), {texts} text shape(s)")
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(str(p), read_only=True, data_only=True)
            try:
                sheets = wb.sheetnames
                rows = 0
                for ws in wb.worksheets:
                    for _row in ws.iter_rows(values_only=True):
                        rows += 1
                        if rows > 5000:
                            break
                charts = sum(len(getattr(ws, "_charts", []) or [])
                             for ws in wb.worksheets)
            finally:
                wb.close()
            if not sheets:
                return _result(p, False, "workbook has no sheets")
            if rows == 0:
                return _result(p, False, "workbook is empty")
            return _result(p, True,
                           f"{len(sheets)} sheet(s), ~{rows} row(s), "
                           f"{charts} chart(s)")
        if suffix in (".png", ".jpg", ".jpeg", ".webp"):
            from PIL import Image
            with Image.open(p) as im:
                im.verify()
            with Image.open(p) as im:
                w, h = im.size
                fmt = im.format
            if w < 8 or h < 8:
                return _result(p, False, f"image is only {w}×{h}")
            return _result(p, True, f"{fmt} {w}×{h}")
        if suffix == ".json":
            json.loads(p.read_text(encoding="utf-8"))
            return _result(p, True, "valid JSON")
        if suffix == ".csv":
            with open(p, newline="", encoding="utf-8", errors="replace") as fh:
                rows = sum(1 for _ in csv.reader(fh))
            if rows < 1:
                return _result(p, False, "CSV has no rows")
            return _result(p, True, f"{rows} row(s)")
        if suffix in (".py",):
            import ast
            ast.parse(p.read_text(encoding="utf-8", errors="replace"))
            return _result(p, True, "Python parses")
        if suffix in (".md", ".txt", ".log"):
            text = p.read_text(encoding="utf-8", errors="replace")
            if len(text.strip()) < 20:
                return _result(p, False, "text file is effectively empty")
            return _result(p, True, f"{len(text)} chars")
        if suffix in (".html", ".htm"):
            text = p.read_text(encoding="utf-8", errors="replace").lower()
            if "<html" not in text and "<body" not in text:
                return _result(p, False, "HTML has no document structure")
            return _result(p, True, f"{len(text)} chars")
        if suffix in (".js", ".jsx", ".ts", ".tsx", ".css", ".svg", ".yaml",
                      ".yml", ".toml", ".zip"):
            return _result(p, True, f"{size} bytes (syntax not parsed)",
                           warning="no parser for this type")
    except Exception as exc:
        return _result(p, False, f"could not be reopened: {exc}")

    return _result(p, True, f"{size} bytes (unrecognised type)",
                   warning="type not recognised; existence only")


def validate_all(artifacts: list[dict], workspace) -> dict:
    """Validate every artifact; a task passes only if all of them open."""
    results = []
    for item in artifacts or []:
        path = item.get("path") if isinstance(item, dict) else item
        results.append(validate_artifact(path, workspace))
    failures = [r for r in results if not r["ok"]]
    return {
        "ok": bool(results) and not failures,
        "results": results,
        "failures": failures,
        "detail": ("; ".join(f"{Path(r['path']).name}: {r['detail']}"
                             for r in failures)
                   if failures else f"{len(results)} artifact(s) verified"),
    }


__all__ = ["validate_artifact", "validate_all"]
