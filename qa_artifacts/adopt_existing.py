"""qa_artifacts/adopt_existing.py — verify what is already on disk and record it.

The batch runner keeps one state file per artifact so a run that dies half way
never loses the artifacts that already succeeded. This script fills that state
from artifacts that were built earlier, by VERIFYING them rather than by
trusting that they exist:

  pages   check_html + browser_audit on the built index.html
  docs    the file opens and has real content (python-docx / openpyxl / pptx)
  apps    the recipe's files are present and the code compiles (py_compile,
          node --check, --selftest where the recipe asks for it)

    python qa_artifacts/adopt_existing.py
"""
from __future__ import annotations

import json
import subprocess
import sys
import zipfile
from pathlib import Path

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE))

from qa_artifacts.build_two_of_each import (  # noqa: E402
    APPS, DOCUMENTS, OUT, PAGES, SELLING, _load_states, _missing_recipe_files,
    _report, _save_state,
)

STATE_DIR = BASE / "qa_artifacts" / "reports" / "two_of_each_state"


def _run(cmd: list[str], cwd: Path, timeout: int = 240) -> tuple[int, str]:
    try:
        proc = subprocess.run(cmd, cwd=str(cwd), capture_output=True, text=True,
                              encoding="utf-8", errors="replace", timeout=timeout)
        return proc.returncode, (proc.stdout + proc.stderr).strip()
    except subprocess.TimeoutExpired:
        return -9, "timed out"
    except Exception as e:                       # never stop the run
        return -1, f"{type(e).__name__}: {e}"


def adopt_pages() -> list[dict]:
    """Verify each built page the way build_site verifies it."""
    from core.site_checks import browser_audit, check_html

    out = []
    for slug, kind, studio, brief in PAGES:
        folder = OUT / "pages" / slug
        index = folder / "index.html"
        if not index.exists():
            continue
        html = index.read_text(encoding="utf-8", errors="replace")
        ok, problems = check_html(index, kind)
        audit = browser_audit(index) if ok else {"status": "skipped"}
        if audit.get("status") == "failed":
            ok = False
            problems = problems + audit.get("problems", [])
        plan = {}
        plan_file = folder / "plan.json"
        if plan_file.exists():
            try:
                plan = json.loads(plan_file.read_text(encoding="utf-8"))
            except Exception:
                plan = {}
        generator = ""
        import re
        m = re.search(r'name="generator" content="([^"]+)"', html)
        if m:
            generator = m.group(1)
        out.append({
            "slug": slug, "artifact": "page", "kind": kind,
            "studio_requested": studio,
            "studio": generator.replace("Crown Code Studio — ", "").strip(),
            "studio_used": generator,
            "path": str(index),
            "bytes": index.stat().st_size,
            "ok": ok, "finished": ok,
            "problems": problems,
            "browser": audit.get("status", "not run"),
            "browser_problems": audit.get("problems", []),
            "checklist_score": "",
            "adopted": True,
            "message": "verified from disk with check_html + browser_audit",
        })
        print(f"  page {slug:<32} ok={ok} browser={audit.get('status')} "
              f"{problems[:1]}", flush=True)
    return out


def _doc_content_ok(path: Path) -> tuple[bool, str]:
    """Open the document the way a user would and look for real content."""
    try:
        if path.suffix == ".docx":
            from docx import Document
            doc = Document(str(path))
            text = "\n".join(p.text for p in doc.paragraphs)
            tables = len(doc.tables)
            return (len(text.strip()) > 200 and tables > 0,
                    f"{len(text.strip())} chars, {tables} tables")
        if path.suffix == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(str(path))
            sheets = wb.sheetnames
            cells = sum(1 for ws in wb.worksheets for row in ws.iter_rows()
                        for c in row if c.value not in (None, ""))
            return (cells > 20, f"{len(sheets)} sheets, {cells} filled cells")
        if path.suffix == ".pptx":
            from pptx import Presentation
            prs = Presentation(str(path))
            text = "\n".join(sh.text_frame.text for sl in prs.slides
                             for sh in sl.shapes if sh.has_text_frame)
            tables = sum(1 for sl in prs.slides for sh in sl.shapes if sh.has_table)
            charts = sum(1 for sl in prs.slides for sh in sl.shapes if sh.has_chart)
            return (len(prs.slides) >= 3 and len(text.strip()) > 150,
                    f"{len(prs.slides)} slides, {len(text.strip())} chars, "
                    f"{tables} tables, {charts} charts")
    except Exception as e:
        return False, f"{type(e).__name__}: {e}"
    return False, "unknown type"


def adopt_docs() -> list[dict]:
    out = []
    for title, instruction, doc_type, detail in DOCUMENTS:
        slug = "doc_" + title.split(" — ")[0].replace(" ", "_")[:32]
        folder = OUT / "documents" / doc_type
        path = folder / f"{slug}.{doc_type}"
        if not path.exists():
            continue
        ok, detail_text = _doc_content_ok(path)
        out.append({
            "slug": slug, "artifact": doc_type, "title": title,
            "path": str(path), "bytes": path.stat().st_size,
            "ok": ok, "finished": ok, "detail": detail_text,
            "adopted": True,
            "message": "verified from disk: " + detail_text,
        })
        print(f"  doc  {slug:<34} ok={ok} {detail_text}", flush=True)
    return out


def adopt_app(slug: str, kind: str, workspace: Path) -> dict | None:
    """An app counts as built only when its recipe files are on disk AND the
    code that can be checked, checks out."""
    if not workspace.exists():
        return None
    missing = _missing_recipe_files(kind, workspace)
    files = sorted(str(p.relative_to(workspace)).replace("\\", "/")
                   for p in workspace.rglob("*")
                   if p.is_file() and "node_modules" not in p.parts)
    checks: list[dict] = []

    py = sorted(p for p in workspace.rglob("*.py")
                if "node_modules" not in p.parts and ".venv" not in p.parts)
    if py:
        code, out = _run([sys.executable, "-m", "py_compile"] +
                         [str(p.relative_to(workspace)) for p in py[:20]],
                         workspace)
        checks.append({"check": "py_compile", "ok": code == 0,
                       "output": out[-300:]})
    js = sorted(p for p in workspace.rglob("*.js")
                if "node_modules" not in p.parts)
    for path in js[:20]:
        code, out = _run(["node", "--check", str(path.relative_to(workspace))],
                         workspace)
        checks.append({"check": f"node --check {path.name}", "ok": code == 0,
                       "output": out[-300:]})
    if kind == "python-desktop":
        code, out = _run([sys.executable, "main.py", "--selftest"], workspace)
        checks.append({"check": "main.py --selftest", "ok": code == 0,
                       "output": out[-300:]})

    code_ok = all(c["ok"] for c in checks) if checks else True
    ok = (not missing) and code_ok
    return {
        "slug": slug, "artifact": kind,
        "files": files[:80], "file_count": len(files),
        "missing_files": missing, "checks": checks,
        "ok": ok, "finished": ok,
        "adopted": True,
        "message": "verified from disk: recipe files present" if not missing
                   else "verified from disk: missing " + ", ".join(missing),
    }


def adopt_apps() -> list[dict]:
    out = []
    for spec in APPS:
        workspace = OUT / "apps" / spec["slug"]
        entry = adopt_app(spec["slug"], spec["kind"], workspace)
        if entry is None:
            continue
        out.append(entry)
        print(f"  app  {spec['slug']:<32} ok={entry['ok']} "
              f"files={entry['file_count']} missing={entry['missing_files']}",
              flush=True)
    for spec in SELLING:
        workspace = OUT / "selling" / spec["slug"]
        entry = adopt_app(spec["slug"], spec["kind"], workspace)
        if entry is None:
            continue
        out.append(entry)
        print(f"  store {spec['slug']:<32} ok={entry['ok']} "
              f"files={entry['file_count']}", flush=True)
    return out


def main() -> int:
    known = {e.get("slug") for e in _load_states()}
    entries = []
    print("pages:", flush=True)
    entries += adopt_pages()
    print("documents:", flush=True)
    entries += adopt_docs()
    print("apps:", flush=True)
    entries += adopt_apps()
    for entry in entries:
        if entry.get("slug") not in known or entry.get("ok"):
            _save_state(entry)
    _report()
    ok = sum(1 for e in entries if e.get("ok"))
    print(f"\nadopted {len(entries)} artifacts · {ok} verified ok")
    for e in entries:
        print(f"  {e.get('slug'):<38} {'ok' if e.get('ok') else 'INCOMPLETE'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
