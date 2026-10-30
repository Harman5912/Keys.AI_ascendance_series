"""End-to-end acceptance tests for the research browser extension.

Each test is a whole workflow the user would actually ask for, run against the
local fixture site (real HTTP, real Chromium, real files on disk).
"""

from __future__ import annotations

import shutil
import sys
import tempfile
import threading
import time
import unittest
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest import mock

from browser_fixtures import ROOT, BrowserCase, stop_site

sys.path.insert(0, str(ROOT))

from PIL import Image  # noqa: E402

from core.research_browser import (collector, integration, library,  # noqa: E402
                                   paths, service as service_mod, tools)


def tearDownModule():
    stop_site()


class _QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass


def serve_folder(folder: Path):
    handler = partial(_QuietHandler, directory=str(folder))
    server = ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    return server


class AcceptanceCase(BrowserCase):
    def setUp(self):
        super().setUp()
        self.temp = Path(tempfile.mkdtemp(prefix="rb_accept_"))
        self._patches = [
            mock.patch.object(paths, "library_path",
                              lambda: self.temp / "library.json"),
            mock.patch.object(paths, "workspace_root",
                              lambda: self.temp / "workspace"),
        ]
        for patch in self._patches:
            patch.start()
        self.task = "accept"

    def tearDown(self):
        for patch in reversed(self._patches):
            patch.stop()
        shutil.rmtree(self.temp, ignore_errors=True)
        super().tearDown()


class TestAcceptance(AcceptanceCase):
    def test_acceptance_1_research_a_real_source(self):
        self.load("/")
        saved = collector.save_source(self.service, self.tab, self.task,
                                      question="What does the page contain?")
        self.assertTrue(saved["ok"], saved)
        source = library.get_source(saved["source"]["source_id"])
        self.assertTrue(source["inspected"])
        self.assertEqual(source["url"], self.site.url("/"))
        self.assertIn("quick brown fox", source["excerpt"])
        self.assertTrue(source["screenshots"])
        self.assertTrue(Path(source["screenshots"][0]).is_file())
        coverage = library.coverage(self.task)
        self.assertEqual(coverage["sources"], 1)
        self.assertEqual(coverage["inspected"], 1)

    def test_acceptance_2_screenshot_a_website(self):
        self.load("/second")
        shot = collector.capture(self.service, self.tab, task_id=self.task,
                                 label="source")
        self.assertTrue(shot["ok"], shot)
        path = Path(shot["path"])
        self.assertTrue(path.is_file())
        with Image.open(path) as im:
            im.load()
            self.assertGreater(im.size[0], 200)
        from PyQt6.QtGui import QPixmap
        self.assertFalse(QPixmap(str(path)).isNull())
        self.assertIn("/second", shot["url"])

    def test_acceptance_3_collect_and_integrate_an_image(self):
        self.load("/")
        collected = collector.collect_image(
            self.service, self.tab, self.site.url("/big.png"),
            task_id=self.task, license_status="open-licence",
            attribution="Fixture, CC0")
        self.assertTrue(collected["ok"], collected)
        asset = collected["asset"]
        library.approve_asset(asset["asset_id"], True)

        site = self.temp / "site"
        (site / "assets").mkdir(parents=True)
        (site / "index.html").write_text(
            "<html><body><h1>Portfolio</h1>\n<!-- KEYS_AI_ASSETS -->\n</body></html>",
            encoding="utf-8")
        result = integration.add_to_web_project(
            library.get_asset(asset["asset_id"])["path"], str(site),
            alt="Collected photograph")
        self.assertTrue(result["ok"], result)
        verified = integration.verify_reference(
            str(site), result["reference"], page=str(site / "index.html"))
        self.assertTrue(verified["ok"], verified)
        # The page must load in the browser and show the image.
        server = serve_folder(site)
        try:
            self.service.navigate(self.tab, f"http://127.0.0.1:{server.server_port}/")
            images = self.service.list_page_images(self.tab)["images"]
            self.assertTrue(any(result["reference"] in i["src"] for i in images))
        finally:
            server.shutdown()

    def test_acceptance_4_collect_research_for_development(self):
        # 1. Inspect a design reference and record it.
        self.load("/")
        source = collector.save_source(self.service, self.tab, self.task,
                                       question="layout reference")["source"]
        # 2. Collect an image from it.
        asset = collector.collect_image(
            self.service, self.tab, self.site.url("/big.png"),
            task_id=self.task, license_status="open-licence",
            source_id=source["source_id"])["asset"]
        # 3. Build a small site that uses it.
        site = self.temp / "portfolio"
        (site / "assets").mkdir(parents=True)
        (site / "index.html").write_text(
            "<!doctype html><html><head><title>Portfolio</title></head>"
            "<body><h1>My work</h1>\n<!-- KEYS_AI_ASSETS -->\n</body></html>",
            encoding="utf-8")
        result = integration.add_to_project(asset["path"], str(site),
                                            kind="web-static")
        self.assertTrue(result["ok"], result)
        # 4. Launch it and verify the real page.
        server = serve_folder(site)
        try:
            url = f"http://127.0.0.1:{server.server_port}/index.html"
            nav = self.service.navigate(self.tab, url)
            self.assertTrue(nav["ok"], nav)
            images = self.service.list_page_images(self.tab)["images"]
            self.assertTrue(images, "the built page shows no image")
            shot = collector.capture(self.service, self.tab, task_id=self.task,
                                     label="built_site")
            self.assertTrue(shot["ok"], shot)
            self.assertTrue(Path(shot["path"]).is_file())
        finally:
            server.shutdown()
        self.assertTrue(library.get_source(source["source_id"])["inspected"])

    def test_acceptance_5_collect_information_into_a_report(self):
        self.load("/")
        self.service._run_js(self.tab, """
            (() => { const el = document.getElementById('paragraph');
              const r = document.createRange(); r.selectNodeContents(el);
              const s = window.getSelection(); s.removeAllRanges();
              s.addRange(r); return true; })()""")
        evidence = collector.collect_selection(
            self.service, self.tab, task_id=self.task, claim="Fox fact")
        self.assertTrue(evidence["ok"], evidence)
        table = collector.collect_table(self.service, self.tab, 0,
                                        task_id=self.task)
        self.assertTrue(table["ok"], table)
        self.assertEqual(table["table"]["headers"], ["Name", "Value"])
        shot = collector.capture(self.service, self.tab, task_id=self.task,
                                 label="evidence")
        from docx import Document
        report = self.temp / "report.docx"
        doc = Document()
        doc.add_heading("Research report", 1)
        doc.add_paragraph("Evidence: " + evidence["text"])
        doc.save(str(report))
        inserted = integration.add_to_docx(shot["path"], str(report),
                                            caption="Source page",
                                            source_note=self.site.url("/"))
        self.assertTrue(inserted["ok"], inserted)
        reopened = Document(str(report))
        text = "\n".join(p.text for p in reopened.paragraphs)
        self.assertIn("quick brown fox", text)
        self.assertEqual(len(reopened.inline_shapes), 1)
        self.assertEqual(library.coverage(self.task)["evidence"], 2)

    def test_acceptance_6_recovery_and_permissions(self):
        self.load("/")
        source = collector.save_source(self.service, self.tab, self.task,
                                       question="recovery")["source"]
        # A consequential action is held until the user releases control.
        self.service.request_human_takeover(self.tab, "sign-in required")
        blocked = tools.execute({"action": "click_element", "tab_id": self.tab,
                                 "selector": "#second-link"})
        self.assertIn("failed", blocked.lower())
        self.assertEqual(self.service.state(self.tab)["url"], self.site.url("/"))
        # A fresh service instance still sees the stored source.
        fresh = service_mod.BrowserService()
        try:
            records = library.list_sources(self.task)
            self.assertTrue(any(r["source_id"] == source["source_id"]
                                for r in records))
        finally:
            fresh.shutdown()
        self.service.release_takeover(self.tab)
        released = self.service.click_element(self.tab, selector="#second-link")
        self.assertTrue(released["ok"], released)
        self.assertIn("/second", self.service.state(self.tab)["url"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
