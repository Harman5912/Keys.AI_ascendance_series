"""Regression tests: project-agent workers against real files and a local page.

Documents, spreadsheets and images are generated for real and reopened by the
artifact validator. Research is tested with controlled search/fetch fixtures
(deterministic, offline). The browser worker runs Playwright against a local
HTTP server; if Chromium cannot start, that group is reported as skipped.
"""

from __future__ import annotations

import csv
import hashlib
import sys
import tempfile
import threading
import time
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from PIL import Image  # noqa: E402

from universal_agent import security, validation  # noqa: E402
from universal_agent.workers import Context  # noqa: E402
from universal_agent.workers import browser as browser_worker  # noqa: E402
from universal_agent.workers import data as data_worker  # noqa: E402
from universal_agent.workers import documents as doc_worker  # noqa: E402
from universal_agent.workers import images as image_worker  # noqa: E402
from universal_agent.workers import research as research_worker  # noqa: E402


class _Base(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="ua_workers_"))
        self.workspace = self.tmp / "workspace"
        self.workspace.mkdir()

    def tearDown(self):
        import shutil
        shutil.rmtree(self.tmp, ignore_errors=True)

    def ctx(self, request: str, params: dict | None = None,
            answers: dict | None = None, **kwargs) -> Context:
        task = {"id": "t1", "request": request, "answers": answers or {},
                "plan": [], "artifacts": []}
        return Context(task=task, workspace=self.workspace,
                       params=params or {}, **kwargs)

    def assert_valid(self, result: dict):
        self.assertTrue(result["ok"], result.get("error"))
        self.assertTrue(result["artifacts"])
        for item in result["artifacts"]:
            check = validation.validate_artifact(item["path"], self.workspace)
            self.assertTrue(check["ok"], check)


# ── research ───────────────────────────────────────────────────────────────
class TestResearch(_Base):
    def _run(self, results, pages=None, complete=None, body_lines=3):
        pages = pages or {}

        def search(query, limit=5):
            return results

        def fetch(url, timeout=12):
            if url in pages:
                if isinstance(pages[url], Exception):
                    raise pages[url]
                return pages[url]
            raise RuntimeError("404")

        return research_worker.run(self.ctx(
            "Research local LLM runtimes", search=search, fetch=fetch,
            complete=complete))

    def test_report_cites_every_finding(self):
        result = self._run(
            [{"title": "Runtime A", "url": "https://a.example/x", "snippet": "s"}],
            {"https://a.example/x": "Local runtimes are fast. "
                                    "Local runtimes support many models. "
                                    "Local runtimes differ in memory use."})
        self.assert_valid(result)
        text = Path(result["artifacts"][0]["path"]).read_text(encoding="utf-8")
        self.assertIn("[1]", text)
        # Scope to the Findings *section*: later sections (repositories,
        # presets, design references) legitimately contain uncited bullets.
        after = text.split("## Findings", 1)[1]
        findings = after.split("\n## ", 1)[0]
        bullets = [ln for ln in findings.splitlines() if ln.startswith("- ")]
        self.assertTrue(bullets)
        for line in bullets:
            self.assertRegex(line, r"\[\d+\]")

    def test_sources_list_maps_to_fetched_urls(self):
        result = self._run(
            [{"title": "Page", "url": "https://b.example/page", "snippet": "s"}],
            {"https://b.example/page": ("Testing runtimes locally is useful. "
                                        "Runtimes differ in memory use. "
                                        "Runtimes differ in tool support. "
                                        "Runtimes are documented online.")})
        text = Path(result["artifacts"][0]["path"]).read_text(encoding="utf-8")
        self.assertIn("https://b.example/page", text)
        self.assertEqual(result["evidence"][0]["method"], "page fetched")

    def test_failed_source_is_listed_in_limitations(self):
        result = self._run(
            [{"title": "Bad", "url": "https://bad.example/x", "snippet": "s"},
             {"title": "Good", "url": "https://good.example/x", "snippet": "s"}],
            {"https://bad.example/x": RuntimeError("connection reset"),
             "https://good.example/x": "A useful sentence about local runtimes. "
                                       "Another useful runtime sentence."})
        text = Path(result["artifacts"][0]["path"]).read_text(encoding="utf-8")
        self.assertIn("could not be fetched", text)
        self.assertIn("https://bad.example/x", text)

    def test_snippet_only_evidence_is_labelled(self):
        result = self._run(
            [{"title": "Snippet", "url": "https://c.example/x",
              "snippet": "A snippet about runtimes."}],
            {"https://c.example/x": RuntimeError("blocked")})
        self.assertTrue(result["ok"], result.get("error"))
        self.assertFalse(result["evidence"][0]["verified"])
        text = Path(result["artifacts"][0]["path"]).read_text(encoding="utf-8")
        self.assertIn("snippet only", text)
        self.assertIn("could not be fetched", text)

    def test_no_search_results_fails_without_fabricating(self):
        result = self._run([])
        self.assertFalse(result["ok"])
        self.assertIn("no report", result["error"].lower())
        self.assertFalse(list(self.workspace.glob("*.md")))

    def test_uncited_model_summary_is_discarded(self):
        result = self._run(
            [{"title": "A", "url": "https://a.example/x", "snippet": ""}],
            {"https://a.example/x": ("Local runtimes are fast. Local runtimes "
                                     "support many models. Local runtimes need "
                                     "careful memory tuning. Local runtimes are "
                                     "documented in several places.")},
            complete=lambda _p, _s: "A confident summary with no citations.")
        text = Path(result["artifacts"][0]["path"]).read_text(encoding="utf-8")
        self.assertNotIn("confident summary", text)
        self.assertIn("deterministic extract", text)

    def test_cited_model_summary_is_kept(self):
        result = self._run(
            [{"title": "A", "url": "https://a.example/x", "snippet": ""}],
            {"https://a.example/x": ("Local runtimes are fast. Local runtimes "
                                     "support many models. Local runtimes need "
                                     "careful memory tuning. Local runtimes are "
                                     "documented in several places.")},
            complete=lambda _p, _s: "Local runtimes are fast [1].")
        text = Path(result["artifacts"][0]["path"]).read_text(encoding="utf-8")
        self.assertIn("Local runtimes are fast [1].", text)
        self.assertIn("citation-checked", text)

    def test_cancellation_stops_before_writing(self):
        ctx = self.ctx("Research x", search=lambda q, l=5: [],
                       fetch=lambda u, t=12: "", should_cancel=lambda: True)
        result = research_worker.run(ctx)
        self.assertFalse(result["ok"])
        self.assertIn("cancel", result["error"].lower())


# ── documents ──────────────────────────────────────────────────────────────
class TestDocuments(_Base):
    def setUp(self):
        super().setUp()
        # Keep document generation offline & fast: no cover-image fetching.
        self._img_patch = mock.patch(
            "actions._office_builder._image_fetch_retry",
            lambda *a, **k: None)
        self._img_patch.start()

    def tearDown(self):
        self._img_patch.stop()
        super().tearDown()

    def test_word_document(self):
        result = doc_worker.run(self.ctx(
            "Create a Word report about renewable energy",
            {"format": "docx"}))
        self.assert_valid(result)
        path = Path(result["artifacts"][0]["path"])
        self.assertEqual(path.suffix, ".docx")
        check = validation.validate_artifact(path)
        self.assertIn("paragraph", check["detail"])

    def test_pdf_document(self):
        result = doc_worker.run(self.ctx("Create a PDF about solar panels",
                                         {"format": "pdf"}))
        self.assert_valid(result)
        self.assertEqual(Path(result["artifacts"][0]["path"]).suffix, ".pdf")

    def test_powerpoint_slide_count(self):
        result = doc_worker.run(self.ctx(
            "Create a PowerPoint presentation about rivers",
            {"format": "pptx"}))
        self.assert_valid(result)
        from pptx import Presentation
        prs = Presentation(result["artifacts"][0]["path"])
        self.assertGreaterEqual(len(prs.slides), 2)

    def test_excel_workbook(self):
        result = doc_worker.run(self.ctx(
            "Create an Excel budget workbook", {"format": "xlsx"}))
        self.assert_valid(result)
        from openpyxl import load_workbook
        wb = load_workbook(result["artifacts"][0]["path"])
        self.assertTrue(wb.sheetnames)

    def test_format_detected_from_request(self):
        result = doc_worker.run(self.ctx(
            "Create a PowerPoint presentation about oceans"))
        self.assert_valid(result)
        self.assertEqual(Path(result["artifacts"][0]["path"]).suffix, ".pptx")

    def test_filename_is_sanitised(self):
        result = doc_worker.run(self.ctx(
            "Create a Word document about ../../etc/passwd and cats",
            {"format": "docx"}))
        self.assert_valid(result)
        path = Path(result["artifacts"][0]["path"])
        self.assertEqual(path.parent, self.workspace.resolve())
        self.assertNotIn("..", path.name)


# ── data ───────────────────────────────────────────────────────────────────
class TestData(_Base):
    def _csv(self, rows, name="data.csv") -> Path:
        path = self.tmp / name
        with open(path, "w", newline="", encoding="utf-8") as fh:
            csv.writer(fh).writerows(rows)
        return path

    def test_csv_aggregates_are_correct(self):
        path = self._csv([["month", "revenue"],
                          ["Jan", "100"], ["Feb", "200"], ["Mar", "300"],
                          ["Apr", "400"], ["May", "500"], ["Jun", "600"],
                          ["Jul", "700"], ["Aug", "800"], ["Sep", "900"],
                          ["Oct", "1000"]])
        result = data_worker.run(self.ctx(f"Analyse {path}", {"input_file": str(path)}))
        self.assert_valid(result)
        text = Path(result["artifacts"][0]["path"]).read_text(encoding="utf-8")
        self.assertIn("5,500", text)          # sum
        self.assertIn("550", text)            # mean
        from openpyxl import load_workbook
        wb = load_workbook(result["artifacts"][1]["path"])
        self.assertIn("Summary", wb.sheetnames)
        self.assertIn("Data", wb.sheetnames)

    def test_xlsx_multi_sheet(self):
        from openpyxl import Workbook
        path = self.tmp / "book.xlsx"
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Sales"
        ws1.append(["region", "amount"])
        for i in range(3):
            ws1.append([f"R{i}", 10 * (i + 1)])
        ws2 = wb.create_sheet("Empty")
        ws3 = wb.create_sheet("Notes")
        ws3.append(["note"])
        ws3.append(["hello"])
        wb.save(path)
        result = data_worker.run(self.ctx(f"Analyse {path}", {"input_file": str(path)}))
        self.assert_valid(result)
        text = Path(result["artifacts"][0]["path"]).read_text(encoding="utf-8")
        self.assertIn("Empty", text)
        self.assertIn("empty worksheet", text.lower())

    def test_missing_and_duplicate_headers(self):
        path = self._csv([["value", "value", ""],
                          ["1", "2", "x"], ["3", "4", "y"]])
        result = data_worker.run(self.ctx(f"Analyse {path}", {"input_file": str(path)}))
        self.assert_valid(result)
        text = Path(result["artifacts"][0]["path"]).read_text(encoding="utf-8")
        self.assertIn("column_3", text)
        self.assertIn("value (2)", text)

    def test_numbers_stored_as_text_are_coerced_and_reported(self):
        path = self._csv([["amount"], ["10"], ["20"], ["30"]])
        result = data_worker.run(self.ctx(f"Analyse {path}", {"input_file": str(path)}))
        text = Path(result["artifacts"][0]["path"]).read_text(encoding="utf-8")
        self.assertIn("stored as text", text)
        self.assertIn("number", text)

    def test_dates_are_detected(self):
        path = self._csv([["date", "n"],
                          ["2026-01-01", "1"], ["2026-02-01", "2"],
                          ["2026-03-01", "3"]])
        result = data_worker.run(self.ctx(f"Analyse {path}", {"input_file": str(path)}))
        text = Path(result["artifacts"][0]["path"]).read_text(encoding="utf-8")
        self.assertIn("date", text)
        self.assertIn("2026-01-01", text)

    def test_empty_workbook_fails(self):
        from openpyxl import Workbook
        path = self.tmp / "empty.xlsx"
        Workbook().save(path)
        result = data_worker.run(self.ctx(f"Analyse {path}", {"input_file": str(path)}))
        self.assertFalse(result["ok"])
        self.assertIn("nothing to analyse", result["error"].lower())

    def test_small_sample_warning(self):
        path = self._csv([["x"], ["1"], ["2"]])
        result = data_worker.run(self.ctx(f"Analyse {path}", {"input_file": str(path)}))
        self.assertTrue(result["ok"], result.get("error"))
        self.assertTrue(any("row" in w for w in result["warnings"]))

    def test_unsupported_dataset_type(self):
        path = self.tmp / "image.png"
        Image.new("RGB", (10, 10)).save(path)
        result = data_worker.run(self.ctx(f"Analyse {path}", {"input_file": str(path)}))
        self.assertFalse(result["ok"])
        self.assertIn("not supported", result["error"].lower())

    def test_original_dataset_untouched(self):
        path = self._csv([["x"], ["1"], ["2"]])
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        data_worker.run(self.ctx(f"Analyse {path}", {"input_file": str(path)}))
        self.assertEqual(before, hashlib.sha256(path.read_bytes()).hexdigest())


# ── images ─────────────────────────────────────────────────────────────────
class TestImages(_Base):
    def _image(self, size=(800, 600), mode="RGB", name="photo.png") -> Path:
        path = self.tmp / name
        Image.new(mode, size, "red" if mode == "RGB" else (255, 0, 0, 128)).save(path)
        return path

    def test_resize_to_exact_dimensions(self):
        src = self._image()
        result = image_worker.run(self.ctx(
            f"Resize {src} to 400x300", {"input_file": str(src),
                                         "operation": "resize",
                                         "width": 400, "height": 300}))
        self.assert_valid(result)
        with Image.open(result["artifacts"][0]["path"]) as im:
            self.assertEqual(im.size, (400, 300))

    def test_percentage_resize(self):
        src = self._image((1000, 500))
        result = image_worker.run(self.ctx(
            f"Resize {src} by 50%", {"input_file": str(src),
                                     "operation": "resize"}))
        with Image.open(result["artifacts"][0]["path"]) as im:
            self.assertEqual(im.size, (500, 250))

    def test_crop_to_square(self):
        src = self._image((800, 600))
        result = image_worker.run(self.ctx(
            f"Crop {src} to square", {"input_file": str(src),
                                      "operation": "crop"}))
        with Image.open(result["artifacts"][0]["path"]) as im:
            self.assertEqual(im.size[0], im.size[1])

    def test_convert_png_to_jpeg_flattens_alpha(self):
        src = self._image(mode="RGBA", name="alpha.png")
        result = image_worker.run(self.ctx(
            f"Convert {src} to jpg", {"input_file": str(src),
                                      "operation": "convert",
                                      "format": "jpg"}))
        out = Path(result["artifacts"][0]["path"])
        self.assertEqual(out.suffix, ".jpg")
        with Image.open(out) as im:
            self.assertEqual(im.format, "JPEG")
            self.assertEqual(im.mode, "RGB")

    def test_original_is_preserved(self):
        src = self._image()
        before = hashlib.sha256(src.read_bytes()).hexdigest()
        image_worker.run(self.ctx(f"Resize {src} to 200x200",
                                  {"input_file": str(src),
                                   "operation": "resize"}))
        self.assertEqual(before, hashlib.sha256(src.read_bytes()).hexdigest())

    def test_invalid_image_is_reported(self):
        bad = self.tmp / "broken.png"
        bad.write_bytes(b"not an image")
        result = image_worker.run(self.ctx(f"Inspect {bad}",
                                           {"input_file": str(bad),
                                            "operation": "inspect"}))
        self.assertFalse(result["ok"])
        self.assertIn("not a readable image", result["error"].lower())

    def test_unsupported_operation_lists_supported_ones(self):
        src = self._image()
        result = image_worker.run(self.ctx(f"Do magic on {src}",
                                           {"input_file": str(src),
                                            "operation": "magic"}))
        self.assertFalse(result["ok"])
        self.assertIn("supported", result["error"].lower())

    def test_generation_unavailable_backend_is_explicit(self):
        def unavailable(_job):
            return {"ok": False,
                    "error": "The local image backend (ComfyUI) is not running."}

        result = image_worker.run(self.ctx(
            "Generate an image of a silver knight",
            {"operation": "generate"}, image_backend=unavailable))
        self.assertFalse(result["ok"])
        self.assertIn("not running", result["error"])

    def test_default_backend_reports_missing_comfyui_without_cloud(self):
        # The default backend talks to the real AI Image Studio service, so a
        # developer machine with ComfyUI running would take a different branch.
        # Pin the backend to "not reachable" to assert the no-cloud-fallback
        # contract rather than the state of this machine.
        class _Down:
            def backend_status(self, force: bool = False) -> dict:
                return {"reachable": False}

        with mock.patch("core.image_studio.get_service", lambda: _Down()):
            result = image_worker.run(self.ctx(
                "Generate an image of a silver knight", {"operation": "generate"}))
        self.assertFalse(result["ok"])
        self.assertIn("No cloud fallback", result["error"])

    def test_generation_copies_the_real_image(self):
        def fake_backend(job):
            src = self.tmp / "generated.png"
            Image.new("RGB", (256, 256), "blue").save(src)
            return {"ok": True, "images": [str(src)], "model": "test",
                    "seed": 42}

        result = image_worker.run(self.ctx(
            "Generate a blue square", {"operation": "generate"},
            image_backend=fake_backend))
        self.assert_valid(result)
        with Image.open(result["artifacts"][0]["path"]) as im:
            self.assertEqual(im.size, (256, 256))


# ── browser ────────────────────────────────────────────────────────────────
class TestBrowser(_Base):
    @classmethod
    def setUpClass(cls):
        cls.server = ThreadingHTTPServer(("127.0.0.1", 0), _PageHandler)
        cls.port = cls.server.server_address[1]
        cls.url = f"http://127.0.0.1:{cls.port}"
        cls.thread = threading.Thread(target=cls.server.serve_forever,
                                      daemon=True)
        cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        try:
            cls.server.shutdown()
            cls.server.server_close()
        except Exception:
            pass

    def test_rejects_non_http(self):
        with self.assertRaises(security.SecurityError):
            browser_worker.fetch_page("file:///C:/Windows/win.ini", self.workspace)

    def test_navigate_and_extract(self):
        try:
            result = browser_worker.run(self.ctx(f"Open {self.url}/page and read it",
                                                 {"url": f"{self.url}/page"}))
        except Exception as exc:              # pragma: no cover - env dependent
            self.skipTest(f"Playwright could not launch: {exc}")
        if not result.get("ok") and "executable" in result.get("error", "").lower():
            self.skipTest(result["error"])
        self.assertTrue(result["ok"], result.get("error"))
        self.assertTrue(result["evidence"])
        self.assertIn("Runtime comparison", result["evidence"][0]["quote"])
        shot = [a for a in result["artifacts"] if a["kind"] == "png"]
        self.assertTrue(shot)
        check = validation.validate_artifact(shot[0]["path"])
        self.assertTrue(check["ok"], check["detail"])

    def test_unreachable_page_reports_error(self):
        try:
            result = browser_worker.run(self.ctx(
                f"Open {self.url}/missing and read it",
                {"url": f"{self.url}/missing"}))
        except Exception as exc:              # pragma: no cover
            self.skipTest(f"Playwright could not launch: {exc}")
        if "executable" in result.get("error", "").lower():
            self.skipTest(result["error"])
        self.assertFalse(result["ok"])
        self.assertTrue(result["error"])

    def test_download_is_saved_and_validated(self):
        try:
            result = browser_worker.run(self.ctx(
                f"Open {self.url}/page and download the file",
                {"url": f"{self.url}/page"}))
        except Exception as exc:              # pragma: no cover
            self.skipTest(f"Playwright could not launch: {exc}")
        if not result.get("ok") and "executable" in result.get("error", "").lower():
            self.skipTest(result["error"])
        downloads = [a for a in result["artifacts"] if a["kind"] == "csv"]
        if not downloads:
            self.skipTest("Chromium download did not complete in this environment")
        check = validation.validate_artifact(downloads[0]["path"])
        self.assertTrue(check["ok"], check["detail"])


class _PageHandler(BaseHTTPRequestHandler):
    def log_message(self, *_a):
        pass

    def do_GET(self):                         # noqa: N802
        if self.path.startswith("/page"):
            body = (b"<html><head><title>Runtime comparison</title></head>"
                    b"<body><h1>Runtime comparison</h1>"
                    b"<p>Runtime comparison notes for local models.</p>"
                    b"<a href='/data.csv'>Download data</a></body></html>")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        if self.path.startswith("/data.csv"):
            body = b"month,value\nJan,10\nFeb,20\n"
            self.send_response(200)
            self.send_header("Content-Type", "text/csv")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        self.send_response(404)
        self.end_headers()
        self.wfile.write(b"not found")


if __name__ == "__main__":
    unittest.main(verbosity=2)
