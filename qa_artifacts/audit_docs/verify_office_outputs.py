"""Open every generated office artifact with an independent parser.

This does not trust the builder's own return value: each file is reopened with
python-docx / python-pptx / openpyxl / pypdf and checked for the structure the
request asked for.

Run:  python qa_artifacts/audit_docs/verify_office_outputs.py
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from actions._office_builder import (build_docx, build_pptx,  # noqa: E402
                                     document_plan)
from actions._office_xlsx import build_xlsx  # noqa: E402

OUT = Path(__file__).resolve().parent / "out"

DOCX_TITLE = "Q3 Inventory Report — Warehouse Operations"
DOCX_INSTRUCTION = (
    "A quarterly inventory report for the warehouse team. Include sections on "
    "stock levels, supplier performance and cost breakdown, with a table of "
    "products and quantities."
)
PPTX_TITLE = "Q3 Warehouse Review — Supplier Performance"
PPTX_INSTRUCTION = (
    "A six slide deck for the operations review: stock levels, supplier "
    "performance, cost breakdown, risks and next steps."
)
XLSX_TITLE = "Q3 Inventory Workbook — Stock and Cost"
XLSX_INSTRUCTION = (
    "A workbook with the inventory dataset, a cost summary sheet and a chart of "
    "units per product."
)

results: list[tuple[str, bool, str]] = []


def check(name: str, ok: bool, detail: str) -> None:
    results.append((name, ok, detail))


def audit_docx(path: Path) -> None:
    from docx import Document

    doc = Document(str(path))
    paras = [p for p in doc.paragraphs if p.text.strip()]
    headings = [p.text.strip() for p in doc.paragraphs
                if p.style.name.startswith("Heading") and p.text.strip()]
    body = "\n".join(p.text for p in paras)
    check("docx exists", path.exists(), f"{path.stat().st_size if path.exists() else 0} bytes")
    check("docx has headings", len(headings) >= 3, f"{len(headings)}: {headings[:6]}")
    check("docx has body text", len(paras) >= 10, f"{len(paras)} non-empty paragraphs")
    check("docx has a table", len(doc.tables) >= 1,
          f"{len(doc.tables)} table(s), {sum(len(t.rows) for t in doc.tables)} rows")
    check("docx has no markdown markers", "**" not in body and "##" not in body,
          "no ** or ## in text")
    check("docx mentions the subject", "inventory" in body.lower(),
          body[:80].replace("\n", " "))


def audit_pptx(path: Path) -> None:
    from pptx import Presentation

    prs = Presentation(str(path))
    titles, empty = [], 0
    for slide in prs.slides:
        text = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text.append(shape.text_frame.text)
        joined = "\n".join(text).strip()
        if not joined:
            empty += 1
        title_shape = slide.shapes.title
        titles.append(title_shape.text.strip() if title_shape is not None and title_shape.text else "")
    notes = sum(1 for s in prs.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
    check("pptx exists", path.exists(), f"{path.stat().st_size if path.exists() else 0} bytes")
    check("pptx slide count", 5 <= len(prs.slides) <= 12, f"{len(prs.slides)} slides")
    check("pptx every slide titled", all(titles),
          f"untitled: {[i for i, t in enumerate(titles) if not t]}")
    check("pptx no empty slides", empty == 0, f"{empty} empty")
    check("pptx distinct titles", len(set(titles)) == len(titles), f"{titles}")
    check("pptx widescreen", prs.slide_width > prs.slide_height,
          f"{prs.slide_width} x {prs.slide_height}")
    check("pptx speaker notes", notes > 0, f"{notes} slide(s) with notes")


def audit_xlsx(path: Path) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(str(path))
    sheets = wb.sheetnames
    rows = sum(ws.max_row for ws in wb.worksheets)
    formulas = [c.coordinate for ws in wb.worksheets for row in ws.iter_rows()
                for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    charts = sum(len(ws._charts) for ws in wb.worksheets)
    check("xlsx exists", path.exists(), f"{path.stat().st_size if path.exists() else 0} bytes")
    check("xlsx has sheets", len(sheets) >= 1, f"{sheets}")
    check("xlsx has data rows", rows >= 5, f"{rows} rows across {len(sheets)} sheet(s)")
    check("xlsx has formulas", len(formulas) > 0, f"{len(formulas)}: {formulas[:5]}")
    check("xlsx has a chart", charts >= 1, f"{charts} chart(s)")


def audit_pdf_builder() -> None:
    """The PDF path that file_processor uses for action=create."""
    from actions import file_processor as fp

    dest = OUT / "inventory_summary.pdf"
    if dest.exists():
        dest.unlink()
    out = fp.file_processor({"action": "create", "file_path": str(dest),
                             "title": "Q3 Inventory Report",
                             "instruction": "A short inventory summary with stock levels "
                                            "and total value of 54,200.00."})
    check("pdf create reports success", "Created PDF" in out, out.splitlines()[0][:90])
    check("pdf exists", dest.exists(),
          f"{dest.stat().st_size if dest.exists() else 0} bytes")
    if not dest.exists():
        return
    head = dest.read_bytes()[:5]
    check("pdf magic bytes", head == b"%PDF-", repr(head))
    from pypdf import PdfReader

    reader = PdfReader(str(dest))
    text = "\n".join(p.extract_text() or "" for p in reader.pages)
    check("pdf page count", len(reader.pages) >= 1, f"{len(reader.pages)} page(s)")
    check("pdf has extractable text", len(text.strip()) > 200, f"{len(text.strip())} chars")
    check("pdf text is the real document", "inventory" in text.lower(),
          text.strip()[:90].replace("\n", " "))


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)

    docx_path = OUT / "q3_inventory_report.docx"
    plan = build_docx(docx_path, DOCX_TITLE, DOCX_INSTRUCTION)
    print(f"[docx] builder plan: {plan.kind}, {len(plan.sections)} sections")
    audit_docx(docx_path)

    pptx_path = OUT / "q3_warehouse_review.pptx"
    build_pptx(pptx_path, PPTX_TITLE, PPTX_INSTRUCTION)
    audit_pptx(pptx_path)

    xlsx_path = OUT / "q3_inventory.xlsx"
    build_xlsx(xlsx_path, XLSX_TITLE, XLSX_INSTRUCTION)
    audit_xlsx(xlsx_path)

    audit_pdf_builder()

    width = max(len(n) for n, _, _ in results)
    failed = 0
    print()
    for name, ok, detail in results:
        print(f"{'PASS' if ok else 'FAIL'}  {name.ljust(width)}  {detail}")
        failed += 0 if ok else 1
    print(f"\n{len(results) - failed}/{len(results)} checks passed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
