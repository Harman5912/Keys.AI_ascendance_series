"""Tests for the deterministic Word / PowerPoint / Excel builders."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

BASE = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BASE))

from actions._office_builder import build_docx, build_pptx, document_plan  # noqa: E402
from actions._office_xlsx import build_xlsx  # noqa: E402


class RequirementEnforcementTest(unittest.TestCase):
    """core.doc_builder must ENFORCE the shape the request asks for.

    Phase-2 defect locked down here: an explicit heading list ("Use exactly
    these section headings, in this order: ...") was parsed as no
    requirement at all, so the model's plan shipped unchecked and documents
    silently lost requested sections (Limitations, Future Improvements) and
    stated figures (54,200.00). 2200 output tokens also truncated bigger
    plans mid-JSON.
    """

    DOCX_REQ = (
        "Use EXACTLY these section headings, in this order: Executive "
        "Summary, Objectives, System Architecture, Database Schema, JSON API "
        "Endpoints, Test Results, Limitations, Future Improvements. "
        "Use exactly these verified figures: total inventory value "
        "54,200.00, 32 unit tests passed."
    )

    def _complete_sections(self, req, with_table=True):
        blocks = [["p", "The total is 54,200.00 and 32 tests passed."]]
        if with_table:
            blocks.append(["table", ["Product", "Qty"],
                           [["Keyboard", "10"]]])
        return [{"heading": h, "blocks": list(blocks)}
                for h in req["headings"]]

    def test_heading_list_without_count_is_parsed(self):
        from core.doc_builder import _parse_requirements

        req = _parse_requirements(self.DOCX_REQ)
        self.assertEqual(req["sections"], 8)
        self.assertIn("Limitations", req["headings"])
        self.assertIn("Future Improvements", req["headings"])
        self.assertIn("54,200.00", req["figures"])

    def test_missing_requested_heading_is_reported(self):
        from core.doc_builder import _parse_requirements, _plan_missed

        req = _parse_requirements(self.DOCX_REQ)
        # The right NUMBER of sections with one renamed: only the
        # heading-by-name check can catch this.
        sections = [{"heading": ("Risks" if h == "Limitations" else h),
                     "blocks": [["p", "text"]]}
                    for h in req["headings"]]
        problem = _plan_missed({"sections": sections}, req, "docx")
        self.assertIn("Limitations", problem)

    def test_dropped_figure_is_reported(self):
        from core.doc_builder import _parse_requirements, _plan_missed

        req = _parse_requirements(self.DOCX_REQ)
        sections = [{"heading": h, "blocks": [["p", "some prose"]]}
                    for h in req["headings"]]
        problem = _plan_missed({"sections": sections}, req, "docx")
        self.assertIn("54,200.00", problem)

    def test_complete_plan_passes(self):
        from core.doc_builder import _parse_requirements, _plan_missed

        req = _parse_requirements(self.DOCX_REQ)
        problem = _plan_missed({"sections": self._complete_sections(req)},
                               req, "docx")
        self.assertEqual(problem, "")

    def test_clean_plan_never_cuts_below_the_requested_count(self):
        """A validated 8-section plan was clipped to _BLOCK_LIMIT (6) at
        render time, silently dropping the sections the nudge loop fought
        for (Limitations, Future Improvements)."""
        from core.doc_builder import _clean_plan, _parse_requirements

        req = _parse_requirements(self.DOCX_REQ)
        sections = [{"heading": h,
                     "blocks": [["p", "The total is 54,200.00; "
                                 "32 tests passed."]]}
                    for h in req["headings"]]
        plan = _clean_plan({"sections": sections}, "Report", req)
        self.assertEqual(len(plan["sections"]), 8)
        headings = [s["heading"] for s in plan["sections"]]
        self.assertIn("Limitations", headings)
        self.assertIn("Future Improvements", headings)

    def test_final_attempt_plan_gets_the_requested_figures(self):
        """After every retry the model still dropped the figures; the plan
        then carries them verbatim instead of shipping a document that
        contradicts its own brief."""
        import json

        from core.doc_builder import _inject_figures, _parse_requirements

        req = _parse_requirements(self.DOCX_REQ)
        sections = [{"heading": h,
                     "blocks": [["p", "Generic prose without any number."]]}
                    for h in req["headings"]]
        raw = {"sections": sections}
        _inject_figures(raw, req)
        self.assertIn("54,200.00", json.dumps(raw))
        first = raw["sections"][0]["blocks"][0]
        self.assertEqual(first[0], "bullets")

    def test_slide_count_parses_and_title_slide_is_dropped(self):
        from core.doc_builder import (_effective_headings,
                                      _parse_requirements)

        req = _parse_requirements(
            "Produce EXACTLY 3 slides, titled: Title, Architecture, "
            "Testing.")
        self.assertEqual(req["slides"], 3)
        self.assertEqual(req["sections"], 2)
        self.assertEqual(_effective_headings(req),
                         ["Architecture", "Testing"])


class OfficeBuilderTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory(prefix="keysai_office_")
        self.dir = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_docx_has_sections_lists_and_no_markdown(self):
        from docx import Document

        path = self.dir / "proposal.docx"
        plan = build_docx(path, "Business Proposal — Northwind Bakery Expansion", "")
        self.assertGreaterEqual(len(plan.sections), 6)
        self.assertGreater(path.stat().st_size, 20_000)
        doc = Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Executive Summary", text)
        self.assertIn("Budget", text)
        self.assertNotIn("**", text)
        styles = {p.style.name for p in doc.paragraphs}
        self.assertIn("List Bullet", styles)
        self.assertEqual(len(doc.tables), 2)

    def test_markdown_docx_renders_bold_and_numbered_lists(self):
        from docx import Document

        path = self.dir / "md.docx"
        plan = document_plan("MD", "## Section\n- **Achieve** 20% growth\n1. First step")
        self.assertTrue(plan)
        from actions._office_builder import build_docx_markdown
        build_docx_markdown(path, "MD", "## Section\n- **Achieve** 20% growth\n1. First step")
        doc = Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("**", text)
        self.assertTrue(any(r.bold for p in doc.paragraphs for r in p.runs))
        self.assertIn("Heading 2", {p.style.name for p in doc.paragraphs})
        self.assertIn("List Number", {p.style.name for p in doc.paragraphs})

    def test_pptx_is_widescreen_with_distinct_titled_slides(self):
        from pptx import Presentation

        path = self.dir / "pitch.pptx"
        build_pptx(path, "Seed Pitch — Lumen Fields", "")
        prs = Presentation(str(path))
        self.assertEqual(prs.slide_width, 12192000)
        self.assertGreaterEqual(len(prs.slides), 6)
        titles = [s.shapes.title.text.strip() for s in prs.slides]
        self.assertTrue(all(titles))
        self.assertEqual(len(set(titles)), len(titles))
        charts = sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart)
        self.assertEqual(charts, 1)
        self.assertGreater(path.stat().st_size, 20_000)

    def test_pptx_honours_requested_slide_count_and_topic_outline(self):
        from pptx import Presentation

        cases = [
            ("Seed Pitch — Lumen Fields", "Create an 8-slide investor deck.", 8, "The ask"),
            ("Platform Architecture — Mosaic Gateway", "Create a 7-slide technical review.", 7, "Context"),
            ("Q3 Analytics Review — Northstar", "Create an 8-slide review.", 8, "Headline metrics"),
        ]
        for title, instruction, want, expected_title in cases:
            path = self.dir / f"deck_{want}_{len(expected_title)}.pptx"
            build_pptx(path, title, instruction)
            prs = Presentation(str(path))
            self.assertEqual(len(prs.slides), want, title)
            titles = [s.shapes.title.text.strip() for s in prs.slides]
            self.assertTrue(all(titles), title)
            self.assertEqual(len(set(titles)), len(titles), title)
            self.assertIn(expected_title, titles, title)

    def test_xlsx_adapts_columns_to_the_request(self):
        from openpyxl import load_workbook

        cases = {
            "Q3 Revenue Review — Northstar": "Month",
            "Programme Budget — Atlas Migration": "Workstream",
            "Inventory Forecast — Fern & Flint": "SKU",
        }
        for title, first_header in cases.items():
            path = self.dir / (first_header + ".xlsx")
            build_xlsx(path, title, title)
            wb = load_workbook(str(path))
            self.assertEqual(wb.sheetnames, ["Dashboard", "Trend", "Data", "Assumptions"])
            self.assertEqual(wb["Dashboard"].cell(row=11, column=2).value, first_header)
            formulas = [
                c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")
            ]
            self.assertGreater(len(formulas), 30)
            charts = [c for ws in wb.worksheets for c in ws._charts]
            self.assertEqual(len(charts), 2)
            self.assertEqual(len(wb["Data"].tables), 1)
            if first_header == "SKU":
                # On hand and reorder point must be the only two plotted series.
                self.assertEqual([len(c.series) for c in charts], [2, 1])
                series_cells = [s.tx.strRef.f for s in charts[0].series]
                self.assertTrue(any(t.endswith("C11") for t in series_cells))
                self.assertTrue(any(t.endswith("F11") for t in series_cells))


if __name__ == "__main__":
    unittest.main(verbosity=2)
